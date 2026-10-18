#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
relie-video-transcript：抖音 / 小红书 / 快手 / B 站 视频逐字稿 + 互动数据

流程：TikHub 解析链接（标题、作者、点赞、评论、收藏、分享、播放等）
      -> 拿到视频或音频地址 -> 阿里云百炼 Paraformer 转写 -> 输出 Markdown / CSV / Excel

用法：
    python transcribe.py "<链接或整段分享文案>" ["<链接2>" ...]
    python transcribe.py --file links.txt                # 每行一条，也支持 .csv
    python transcribe.py --file links.txt --limit 5      # 先试跑 5 条
    python transcribe.py ... --stats-only                # 只拿数据，不转写
    python transcribe.py ... --comments 20               # 额外抓前 20 条评论（每页额外计费）
    python transcribe.py ... --douyin-plays              # 额外查抖音播放量（每条额外计费 1 次）
    python transcribe.py --retry-failed                  # 重跑失败的条目
    python transcribe.py                                 # 不带链接：继续处理上次没做完的
    python transcribe.py --export                        # 只重新导出汇总表
    python transcribe.py --check                         # 检查配置和依赖（不调用接口）
任何命令都可以加 --json：结果以 JSON 打印到标准输出，进度日志在标准错误。
"""
import argparse
import csv
import hashlib
import json
import os
import re
import sqlite3
import sys
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import parse_qs, unquote, urlparse

for _stream in (sys.stdout, sys.stderr):
    try:  # Windows 控制台默认 GBK，标题里的表情符号会让 print 崩溃
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

try:
    import requests
except ImportError:
    sys.exit("缺少依赖 requests：请先运行  python -m pip install requests openpyxl")


# ----------------------------------------------------------------------------
# 配置
# 读取顺序（先读到的优先）：系统环境变量 > --config 指定的文件 > ~/.relie/transcript.env > 当前目录 .env
# Key 不要放进技能目录或代码仓库。
# ----------------------------------------------------------------------------
CONFIG_DIR = Path(os.getenv("RELIE_CONFIG_DIR") or (Path.home() / ".relie"))
DEFAULT_CONFIG = CONFIG_DIR / "transcript.env"
cfg = SimpleNamespace()


def load_env(path: Path):
    """读取 .env 格式文件（不依赖第三方库）。已存在的环境变量优先。"""
    if not path or not path.is_file():
        return False
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    return True


def init_config(config_path=None, out_dir=None):
    cfg.loaded = [str(p) for p in (Path(config_path) if config_path else None, DEFAULT_CONFIG, Path.cwd() / ".env")
                  if p and load_env(p)]
    cfg.TIKHUB_API_KEY = os.getenv("TIKHUB_API_KEY", "")
    cfg.TIKHUB_BASE_URL = os.getenv("TIKHUB_BASE_URL", "https://api.tikhub.io").rstrip("/")
    cfg.DASHSCOPE_API_KEY = os.getenv("DASHSCOPE_API_KEY", "")
    cfg.DASHSCOPE_BASE_URL = os.getenv("DASHSCOPE_BASE_URL", "https://dashscope.aliyuncs.com").rstrip("/")
    cfg.ASR_MODEL = os.getenv("ASR_MODEL", "paraformer-v2")
    cfg.PARSE_WORKERS = int(os.getenv("PARSE_WORKERS", "5"))           # TikHub 并发数（官方限速 10 QPS）
    cfg.BATCH_SIZE = min(int(os.getenv("BATCH_SIZE", "100")), 100)      # 每个转写任务的文件数（百炼上限 100）
    cfg.MAX_INFLIGHT_TASKS = int(os.getenv("MAX_INFLIGHT_TASKS", "10"))  # 同时进行中的转写任务数
    cfg.POLL_INTERVAL = int(os.getenv("POLL_INTERVAL", "5"))           # 查询转写进度的间隔（秒）
    # 转写任务超过这么多分钟未完成，就把其中的文件重新提交一次；第二次仍超时才记为失败
    cfg.TASK_TIMEOUT_MIN = int(os.getenv("TASK_TIMEOUT_MIN", "20"))
    cfg.HOTWORDS_VOCAB_ID = os.getenv("HOTWORDS_VOCAB_ID", "")
    cfg.DEBUG_RAW = bool(os.getenv("DEBUG_RAW"))
    cfg.OUT_DIR = Path(out_dir or os.getenv("OUT_DIR") or (Path.cwd() / "relie-transcripts")).resolve()
    cfg.DB_PATH = cfg.OUT_DIR / "state.db"


def log(msg):
    print(f"[{datetime.now():%H:%M:%S}] {msg}", file=sys.stderr, flush=True)


# ----------------------------------------------------------------------------
# 状态库（断点续跑的关键）
# status: pending -> parsed（已拿到数据）-> submitted（转写中）-> done / failed
# ----------------------------------------------------------------------------
COLUMNS = {
    "url": "TEXT PRIMARY KEY", "platform": "TEXT", "raw_input": "TEXT",
    "status": "TEXT DEFAULT 'pending'",
    "item_id": "TEXT", "title": "TEXT", "author": "TEXT", "author_id": "TEXT",
    "publish_time": "TEXT", "tags": "TEXT", "duration_sec": "REAL",
    "plays": "INTEGER", "likes": "INTEGER", "comments": "INTEGER", "collects": "INTEGER",
    "shares": "INTEGER", "coins": "INTEGER", "danmaku": "INTEGER", "stats_at": "TEXT",
    "media_url": "TEXT", "asr_url": "TEXT", "force_upload": "INTEGER DEFAULT 0", "extra": "TEXT",
    "task_id": "TEXT", "submitted_at": "TEXT", "asr_tries": "INTEGER DEFAULT 0",
    "text": "TEXT", "sentences": "TEXT", "top_comments": "TEXT", "md_path": "TEXT",
    "error": "TEXT", "attempts": "INTEGER DEFAULT 0", "updated_at": "TEXT",
}
STATUS_ZH = {"pending": "待处理", "parsed": "已获取数据（未转写）", "submitted": "转写中",
             "done": "完成", "failed": "失败"}
PLATFORM_ZH = {"douyin": "抖音", "xhs": "小红书", "kuaishou": "快手", "bilibili": "B站"}


def db():
    cfg.OUT_DIR.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(cfg.DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("CREATE TABLE IF NOT EXISTS items (%s)" % ", ".join(f"{k} {v}" for k, v in COLUMNS.items()))
    have = {r[1] for r in conn.execute("PRAGMA table_info(items)")}
    for k, v in COLUMNS.items():
        if k not in have:  # 兼容旧版本的状态库
            conn.execute(f"ALTER TABLE items ADD COLUMN {k} {v.replace(' PRIMARY KEY', '')}")
    conn.commit()
    return conn


def upsert(conn, url, **fields):
    fields["updated_at"] = now()
    cols = ", ".join(f"{k}=?" for k in fields)
    conn.execute(f"UPDATE items SET {cols} WHERE url=?", (*fields.values(), url))
    conn.commit()


def get_row(conn, url):
    return conn.execute("SELECT * FROM items WHERE url=?", (url,)).fetchone()


def now():
    return datetime.now().isoformat(timespec="seconds")


# ----------------------------------------------------------------------------
# 读取链接
# ----------------------------------------------------------------------------
URL_RE = re.compile(r"https?://[^\s　，。！？、；：“”‘’（）()<>\"'【】]+")
PLATFORM_HOSTS = {
    "douyin": ("douyin.com", "iesdouyin.com"),
    "xhs": ("xiaohongshu.com", "xhslink.com", "xhslink.cn"),
    "kuaishou": ("kuaishou.com", "kuaishouapp.com", "chenzhongtech.com", "gifshow.com"),
    "bilibili": ("bilibili.com", "b23.tv", "bili2233.cn", "bili22.cn", "bili33.cn"),
}


def detect_platform(url):
    host = (urlparse(url).hostname or "").lower()
    for plat, hosts in PLATFORM_HOSTS.items():
        if any(host == h or host.endswith("." + h) for h in hosts):
            return plat
    return None


def extract_links(lines):
    """每行取第一个支持的平台链接（可以是整段分享文案），按出现顺序去重。"""
    out, seen = [], set()
    for raw in lines:
        for m in URL_RE.finditer(raw):
            url = m.group(0).rstrip(".,;!?")
            plat = detect_platform(url)
            if not plat:
                continue
            if url not in seen:
                seen.add(url)
                out.append((url, plat, raw.strip()))
            break
    return out


def read_file_lines(path: Path):
    text = path.read_text(encoding="utf-8-sig", errors="ignore")
    if path.suffix.lower() == ".csv":
        return [" ".join(r) for r in csv.reader(text.splitlines())]
    return text.splitlines()


# ----------------------------------------------------------------------------
# HTTP 基础
# ----------------------------------------------------------------------------
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
_local = threading.local()


def http():
    """每个线程一个 Session（requests.Session 不保证线程安全）。"""
    s = getattr(_local, "session", None)
    if s is None:
        s = _local.session = requests.Session()
    return s


def tikhub_get(path, params, retries=3):
    url = cfg.TIKHUB_BASE_URL + path
    headers = {"Authorization": f"Bearer {cfg.TIKHUB_API_KEY}"}
    last = None
    for i in range(retries):
        try:
            r = http().get(url, params=params, headers=headers, timeout=60)
            if r.status_code in (429, 500, 502, 503, 504):
                last = f"HTTP {r.status_code}"
                time.sleep(2 * (i + 1))
                continue
            if r.status_code != 200:
                raise RuntimeError(f"TikHub HTTP {r.status_code}: {r.text[:200]}")
            j = r.json()
            if cfg.DEBUG_RAW:
                raw_dir = cfg.OUT_DIR / "raw"
                raw_dir.mkdir(parents=True, exist_ok=True)
                name = f"{path.strip('/').replace('/', '_')}_{int(time.time() * 1000)}.json"
                (raw_dir / name).write_text(json.dumps({"params": params, "resp": j}, ensure_ascii=False, indent=1),
                                            encoding="utf-8")
            if j.get("code") not in (200, None):
                raise RuntimeError(f"TikHub 返回错误: {str(j)[:200]}")
            return j.get("data") or {}
        except requests.RequestException as e:
            last = str(e)
            time.sleep(2 * (i + 1))
    raise RuntimeError(f"TikHub 请求失败: {last}")


# ----------------------------------------------------------------------------
# 通用取值工具（应对接口字段层级变化）
# ----------------------------------------------------------------------------
def find_key(obj, keys, depth=0):
    """在嵌套 JSON 里找第一个出现的 key（值为 0 也算找到）。"""
    if depth > 8:
        return None
    if isinstance(obj, dict):
        for k in keys:
            if k in obj and obj[k] not in (None, "", [], {}):
                return obj[k]
        for v in obj.values():
            r = find_key(v, keys, depth + 1)
            if r is not None:
                return r
    elif isinstance(obj, list):
        for v in obj:
            r = find_key(v, keys, depth + 1)
            if r is not None:
                return r
    return None


def find_dict_with(obj, keys, depth=0):
    """找第一个同时包含所有 keys 的字典。"""
    if depth > 8:
        return None
    if isinstance(obj, dict):
        if all(k in obj for k in keys):
            return obj
        children = obj.values()
    elif isinstance(obj, list):
        children = obj
    else:
        return None
    for v in children:
        r = find_dict_with(v, keys, depth + 1)
        if r is not None:
            return r
    return None


def first(d, keys):
    """在一个字典的同一层里按顺序取第一个非空值。"""
    if not isinstance(d, dict):
        return None
    for k in keys:
        if d.get(k) not in (None, ""):
            return d[k]
    return None


def to_count(v):
    """把 1234 / "1234" / "1.2万" / "3w" / "10万+" / "1,234" 统一成整数。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip().replace(",", "").rstrip("+")
    mult = 1
    for suf, m in (("亿", 10 ** 8), ("万", 10 ** 4), ("w", 10 ** 4), ("W", 10 ** 4), ("k", 10 ** 3), ("K", 10 ** 3)):
        if s.endswith(suf):
            s, mult = s[:-len(suf)], m
            break
    try:
        return int(float(s) * mult)
    except ValueError:
        return None


def fmt_time(v):
    """秒或毫秒时间戳 -> 'YYYY-MM-DD HH:MM'；已经是文字的原样返回。"""
    if v in (None, ""):
        return None
    try:
        t = float(v)
    except (TypeError, ValueError):
        return str(v)
    if t <= 0:
        return None
    if t > 1e12:
        t /= 1000
    return datetime.fromtimestamp(t).strftime("%Y-%m-%d %H:%M")


def to_seconds(v):
    """时长：大于 1000 视为毫秒。只作初值，转写完成后会用音频实际时长覆盖。"""
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    if v <= 0:
        return None
    return round(v / 1000, 1) if v > 1000 else v


HASHTAG_RE = re.compile(r"#([^\s#\[\]]+)")


def tags_text(names, title=""):
    names = [str(n).strip().strip("#") for n in (names or []) if n]
    if not names:
        names = HASHTAG_RE.findall(title or "")
    return " ".join("#" + n for n in dict.fromkeys(names) if n) or None


def find_video_urls(obj, depth=0):
    """兜底：递归收集所有看起来像视频文件的链接。"""
    found = []
    if depth > 10:
        return found
    if isinstance(obj, dict):
        for v in obj.values():
            found += find_video_urls(v, depth + 1)
    elif isinstance(obj, list):
        for v in obj:
            found += find_video_urls(v, depth + 1)
    elif isinstance(obj, str) and obj.startswith("http"):
        s = obj.lower()
        if ".mp4" in s or "sns-video" in s or "douyinvod" in s or "/aweme/v1/play" in s or ("video" in s and "xhscdn" in s):
            if not any(x in s for x in (".jpg", ".jpeg", ".png", ".webp", "/img", "avatar", "cover")):
                found.append(obj)
    return found


def pick_best(urls):
    """优先用 CDN 直链，其次是跳转链接；去掉重复。"""
    uniq = [u for u in dict.fromkeys(urls) if isinstance(u, str) and u.startswith("http")]
    direct = [u for u in uniq if "/aweme/v1/play" not in u and "playwm" not in u]
    return (direct or uniq or [None])[0]


# ----------------------------------------------------------------------------
# 各平台解析：返回作品信息 + 互动数据 + 媒体地址
# ----------------------------------------------------------------------------
def parse_douyin(url, want_media):
    data, errors = None, []
    for path in ("/api/v1/douyin/web/fetch_one_video_by_share_url",
                 "/api/v1/douyin/app/v3/fetch_one_video_by_share_url"):
        try:
            data = tikhub_get(path, {"share_url": url})
            if find_key(data, ["aweme_detail"]):
                break
        except Exception as e:
            errors.append(str(e))
    aweme = find_key(data or {}, ["aweme_detail"])
    if not isinstance(aweme, dict):
        if not data:
            raise RuntimeError("; ".join(errors) or "解析结果为空")
        reason = find_key(data, ["filter_detail", "filter_list"])
        raise RuntimeError(f"作品不可用（可能已删除/私密/仅部分人可见）：{str(reason)[:150]}")

    video = aweme.get("video") or {}
    urls = []
    for br in video.get("bit_rate") or []:
        urls += (br.get("play_addr") or {}).get("url_list") or []
    urls += (video.get("play_addr") or {}).get("url_list") or []
    vurl = pick_best(urls or find_video_urls(video))
    if want_media and not vurl:
        raise RuntimeError("没有找到视频地址（可能是图文作品）")

    stats = aweme.get("statistics") or {}
    author = aweme.get("author") or {}
    names = [x.get("hashtag_name") for x in aweme.get("text_extra") or [] if isinstance(x, dict)]
    names += [x.get("cha_name") for x in aweme.get("cha_list") or [] if isinstance(x, dict)]
    title = (aweme.get("desc") or "").strip()
    return {
        "item_id": str(aweme.get("aweme_id") or ""),
        "title": title,
        "author": author.get("nickname", ""),
        "author_id": str(first(author, ["unique_id", "short_id", "sec_uid"]) or ""),
        "publish_time": fmt_time(aweme.get("create_time")),
        "tags": tags_text(names, title),
        "duration_sec": to_seconds(video.get("duration") or aweme.get("duration")),
        # 抖音大多数接口已不返回播放数（返回 0），需要 --douyin-plays 单独查询
        "plays": to_count(stats.get("play_count")) or None,
        "likes": to_count(stats.get("digg_count")),
        "comments": to_count(stats.get("comment_count")),
        "collects": to_count(stats.get("collect_count")),
        "shares": to_count(stats.get("share_count")),
        "media_url": vurl,
    }


def douyin_plays(item_id):
    data = tikhub_get("/api/v1/douyin/app/v3/fetch_video_statistics", {"aweme_ids": item_id})
    return to_count(find_key(data, ["play_count"]))


def xhs_source_note(data):
    """小红书接口返回的是一个信息流：第一条（source_note）才是请求的笔记，后面是推荐的其他笔记。
    只能从第一条取数据，否则笔记失效时会拿到别人的视频。"""
    notes = data.get("data") if isinstance(data, dict) else None
    if not isinstance(notes, list):
        notes = find_key(data, ["note_list", "items"])
    notes = [n for n in notes or [] if isinstance(n, dict)]
    if not notes:
        msg = find_key(data, ["msg", "message"]) or ""
        raise RuntimeError(f"解析结果里没有笔记（链接可能失效）{msg}")
    note = next((n for n in notes if n.get("source_note")), notes[0])
    if note.get("model_type") == "error":
        raise RuntimeError(f"作品不可用：{note.get('text') or '笔记已删除、不可见或不是视频笔记'}")
    return note


def parse_xhs(url, want_media):
    # TikHub 对小红书无效链接也会扣费，所以只请求一次视频笔记接口
    note = xhs_source_note(tikhub_get("/api/v1/xiaohongshu/app_v2/get_video_note_detail", {"share_text": url}))
    urls = []
    for codec in ("h264", "h265", "av1"):  # h264 兼容性最好
        for s in find_key(note, [codec]) or []:
            if isinstance(s, dict):
                urls += [s.get("master_url")] + list(s.get("backup_urls") or [])
    urls = [u for u in urls if isinstance(u, str) and u.startswith("http")]
    vurl = pick_best(urls or find_video_urls(note))
    if not vurl:
        raise RuntimeError("没有找到视频地址（可能是图文笔记）")

    user = note.get("user") or {}
    title = (note.get("title") or (note.get("desc") or "").split("\n")[0]).strip()
    video = note.get("video_info_v2") or note.get("video") or {}
    return {
        "item_id": str(note.get("id") or note.get("note_id") or ""),
        "title": title,
        "author": first(user, ["nickname", "name"]) or "",
        "author_id": str(first(user, ["red_id", "userid", "id"]) or ""),
        "publish_time": fmt_time(note.get("time")),
        "tags": tags_text([t.get("name") for t in note.get("hash_tag") or [] if isinstance(t, dict)], title),
        "duration_sec": to_seconds(find_key(video, ["duration"])),
        "plays": None,  # 小红书不公开播放量（view_count 恒为 0）
        "likes": to_count(note.get("liked_count")),
        "comments": to_count(note.get("comments_count")),
        "collects": to_count(note.get("collected_count")),
        "shares": to_count(note.get("shared_count")),
        "media_url": vurl,
    }


def kuaishou_video_urls(data):
    urls = []
    for key in ("main_mv_urls", "mainMvUrls"):
        for x in find_key(data, [key]) or []:
            if isinstance(x, dict) and x.get("url"):
                urls.append(x["url"])
    for key in ("photoUrl", "photo_url", "mp4_url", "mp4Url"):
        v = find_key(data, [key])
        if isinstance(v, str):
            urls.append(v)
    return urls + find_video_urls(data)


def parse_kuaishou(url, want_media):
    data, errors = None, []
    for path in ("/api/v1/kuaishou/app/fetch_one_video_by_url", "/api/v1/kuaishou/web/fetch_one_video"):
        try:
            d = tikhub_get(path, {"share_text": url})
            data = d
            if not want_media or kuaishou_video_urls(d):
                break
        except Exception as e:
            errors.append(str(e))
    if not data:
        raise RuntimeError("; ".join(errors) or "解析结果为空")
    vurl = pick_best(kuaishou_video_urls(data))
    if want_media and not vurl:
        raise RuntimeError("没有找到视频地址（可能是图集作品或作品已删除）")

    photo = find_key(data, ["photo", "visionVideoDetail"])
    scope = photo if isinstance(photo, dict) else data
    author = find_key(data, ["author"])
    title = (find_key(scope, ["caption", "title"]) or "").strip()
    tag_list = find_key(scope, ["tags", "tag_list", "tagList"]) or []
    return {
        "item_id": str(find_key(scope, ["photo_id", "photoId", "id"]) or ""),
        "title": title,
        "author": (first(author, ["name", "user_name"]) if isinstance(author, dict) else None)
                  or find_key(data, ["user_name", "userName", "author_name"]) or "",
        "author_id": str((first(author, ["id", "user_id"]) if isinstance(author, dict) else None)
                         or find_key(data, ["user_id", "userId"]) or ""),
        "publish_time": fmt_time(find_key(scope, ["timestamp", "create_time", "time"])),
        "tags": tags_text([first(t, ["name", "tag_name", "tagName"]) for t in tag_list if isinstance(t, dict)], title),
        "duration_sec": to_seconds(find_key(scope, ["duration"])),
        "plays": to_count(find_key(scope, ["view_count", "viewCount"])),
        "likes": to_count(find_key(scope, ["like_count", "likeCount", "realLikeCount"])),
        "comments": to_count(find_key(scope, ["comment_count", "commentCount"])),
        "collects": to_count(find_key(scope, ["collect_count", "collectCount"])),
        "shares": to_count(find_key(scope, ["share_count", "shareCount", "forward_count", "forwardCount"])),
        "media_url": vurl,
    }


BV_RE = re.compile(r"BV[0-9A-Za-z]{10}")


def resolve_bilibili(url):
    """拿到 BV 号和分 P 序号；b23.tv 等短链在本地跟随跳转（不经过 TikHub，不计费）。"""
    m = BV_RE.search(url)
    if not m and detect_platform(url) == "bilibili" and "bilibili.com" not in url:
        r = http().get(url, headers={"User-Agent": UA}, allow_redirects=True, timeout=20, stream=True)
        r.close()
        url = r.url
        m = BV_RE.search(url)
    if not m:
        raise RuntimeError("没有识别到 B 站视频的 BV 号（暂不支持番剧、直播、av 号链接）")
    try:
        p = int(parse_qs(urlparse(url).query).get("p", ["1"])[0])
    except ValueError:
        p = 1
    return m.group(0), max(p, 1)


def parse_bilibili(url, want_media):
    bvid, p = resolve_bilibili(url)
    data = tikhub_get("/api/v1/bilibili/web/fetch_one_video", {"bv_id": bvid})
    view = find_dict_with(data, ("bvid", "stat"))
    if not view:
        msg = find_key(data, ["message", "msg"]) or ""
        raise RuntimeError(f"作品不可用（可能已删除或需要登录）{msg}")
    stat = view.get("stat") or {}
    pages = view.get("pages") or []
    page = pages[p - 1] if len(pages) >= p else (pages[0] if pages else {})
    title = (view.get("title") or "").strip()
    if len(pages) > 1:
        title += f"（P{page.get('page', p)} {page.get('part', '')}）".replace(" ）", "）")
    owner = view.get("owner") or {}
    return {
        "item_id": bvid if p == 1 else f"{bvid}_p{p}",
        "title": title,
        "author": owner.get("name", ""),
        "author_id": str(owner.get("mid") or ""),
        "publish_time": fmt_time(view.get("pubdate")),
        "tags": tags_text([view.get("tname")] if view.get("tname") else [], view.get("desc") or ""),
        "duration_sec": to_seconds(page.get("duration") or view.get("duration")),
        "plays": to_count(stat.get("view")),
        "likes": to_count(stat.get("like")),
        "comments": to_count(stat.get("reply")),
        "collects": to_count(stat.get("favorite")),
        "shares": to_count(stat.get("share")),
        "coins": to_count(stat.get("coin")),
        "danmaku": to_count(stat.get("danmaku")),
        # B 站音视频地址有防盗链，百炼无法直接下载；转写前再取音频地址、下载后上传
        "media_url": None,
        "extra": {"bvid": bvid, "cid": page.get("cid") or view.get("cid"), "aid": view.get("aid")},
    }


def bilibili_audio_urls(bvid, cid):
    """返回候选音频地址（主地址 + 备用镜像），下载时依次尝试。"""
    data = tikhub_get("/api/v1/bilibili/web/fetch_video_playurl", {"bv_id": bvid, "cid": cid})
    audios = [a for a in (find_key(data, ["audio"]) or []) if isinstance(a, dict)]
    audios.sort(key=lambda a: a.get("bandwidth") or 0)  # 语音识别用最低码率就够，下载最快
    for a in audios:
        urls = [first(a, ["baseUrl", "base_url"])] + list(first(a, ["backupUrl", "backup_url"]) or [])
        urls = [u for u in urls if u]
        if urls:
            return urls
    durl = find_key(data, ["durl"]) or []
    if durl and isinstance(durl[0], dict) and durl[0].get("url"):
        return [durl[0]["url"]] + list(durl[0].get("backup_url") or [])
    raise RuntimeError("没有拿到 B 站音频地址（可能是付费/会员专享视频）")


PARSERS = {"douyin": parse_douyin, "xhs": parse_xhs, "kuaishou": parse_kuaishou, "bilibili": parse_bilibili}


# ----------------------------------------------------------------------------
# 评论（可选，按页额外计费）
# ----------------------------------------------------------------------------
def norm_comment(c):
    text = first(c, ["text", "content", "message"]) or ""
    if isinstance(text, dict):  # B 站：content.message
        text = text.get("message") or ""
    user = first(c, ["user", "member", "author", "user_info"]) or {}
    name = first(user, ["nickname", "uname", "name", "user_name"]) if isinstance(user, dict) else None
    name = name or first(c, ["author_name", "authorName", "user_name", "nickname"]) or ""
    return {
        "user": str(name),
        "text": str(text).strip(),
        "likes": to_count(first(c, ["digg_count", "like_count", "likedCount", "liked_count", "like", "likeCount"])),
        "time": fmt_time(first(c, ["create_time", "ctime", "time", "timestamp"])),
    }


def fetch_comments(row, n):
    plat, extra = row["platform"], json.loads(row["extra"] or "{}")
    out, cursor, index, area = [], None, 0, "UNFOLDED"
    max_pages = (n + 19) // 20 + 1
    for page in range(1, max_pages + 1):
        if plat == "douyin":
            d = tikhub_get("/api/v1/douyin/web/fetch_video_comments",
                           {"aweme_id": row["item_id"], "cursor": cursor or 0, "count": 20})
            items = find_key(d, ["comments"]) or []
            cursor, more = find_key(d, ["cursor"]), find_key(d, ["has_more"])
        elif plat == "xhs":
            params = {"sort_strategy": "like_count", "cursor": cursor or "", "index": index, "pageArea": area}
            params.update({"note_id": row["item_id"]} if row["item_id"] else {"share_text": row["url"]})
            d = tikhub_get("/api/v1/xiaohongshu/app_v2/get_note_comments", params)
            items = find_key(d, ["comments", "comment_list"]) or []
            cursor, more = find_key(d, ["cursor"]), find_key(d, ["has_more"])
            index, area = find_key(d, ["index"]) or 0, find_key(d, ["pageArea", "page_area"]) or "UNFOLDED"
        elif plat == "kuaishou":
            d = tikhub_get("/api/v1/kuaishou/app/fetch_video_comment", {"photo_id": row["item_id"], "pcursor": cursor or ""})
            items = find_key(d, ["rootComments", "root_comments", "commentList", "comments"]) or []
            cursor = find_key(d, ["pcursor"])
            more = cursor not in (None, "", "no_more")
        elif plat == "bilibili":
            d = tikhub_get("/api/v1/bilibili/web/fetch_video_comments", {"bv_id": extra.get("bvid"), "pn": page})
            items = find_key(d, ["replies"]) or []
            more = bool(items)
        else:
            return []
        out += [norm_comment(c) for c in items if isinstance(c, dict)]
        if len(out) >= n or not items or not more:
            break
    return [c for c in out if c["text"]][:n]


# ----------------------------------------------------------------------------
# 百炼：上传（防盗链平台用）+ 转写
# ----------------------------------------------------------------------------
REFERERS = {"douyin": "https://www.douyin.com/", "xhs": "https://www.xiaohongshu.com/",
            "kuaishou": "https://www.kuaishou.com/", "bilibili": "https://www.bilibili.com/"}


def ds_headers(async_=False, oss=False):
    h = {"Authorization": f"Bearer {cfg.DASHSCOPE_API_KEY}", "Content-Type": "application/json"}
    if async_:
        h["X-DashScope-Async"] = "enable"
    if oss:
        h["X-DashScope-OssResourceResolve"] = "enable"
    return h


def download(urls, platform, dest: Path):
    """依次尝试候选地址。B 站 CDN 不带 Range 头会返回 567，所以统一带上 Range: bytes=0-。"""
    referer = REFERERS.get(platform, "")
    headers = {"User-Agent": UA, "Referer": referer, "Origin": referer.rstrip("/"), "Range": "bytes=0-"}
    errors = []
    for url in [urls] if isinstance(urls, str) else urls:
        try:
            with http().get(url, headers=headers, stream=True, timeout=120) as r:
                if r.status_code not in (200, 206):
                    errors.append(f"HTTP {r.status_code}")
                    continue
                with open(dest, "wb") as f:
                    for chunk in r.iter_content(1 << 20):
                        f.write(chunk)
            return dest
        except requests.RequestException as e:
            errors.append(str(e)[:100])
    raise RuntimeError(f"下载媒体文件失败：{'；'.join(errors)}")


def upload_to_dashscope(path: Path):
    """上传到百炼临时存储（48 小时有效），返回 oss:// 地址。"""
    r = http().get(f"{cfg.DASHSCOPE_BASE_URL}/api/v1/uploads",
                   params={"action": "getPolicy", "model": cfg.ASR_MODEL},
                   headers={"Authorization": f"Bearer {cfg.DASHSCOPE_API_KEY}"}, timeout=60)
    if r.status_code != 200:
        raise RuntimeError(f"获取百炼上传凭证失败 HTTP {r.status_code}: {r.text[:200]}")
    pol = r.json()["data"]
    size_mb = path.stat().st_size / 1024 / 1024
    if pol.get("max_file_size_mb") and size_mb > float(pol["max_file_size_mb"]):
        raise RuntimeError(f"文件 {size_mb:.0f}MB 超过百炼临时存储上限 {pol['max_file_size_mb']}MB")
    key = f"{pol['upload_dir']}/{path.name}"
    fields = {
        "OSSAccessKeyId": pol["oss_access_key_id"], "Signature": pol["signature"], "policy": pol["policy"],
        "x-oss-object-acl": pol["x_oss_object_acl"], "x-oss-forbid-overwrite": pol["x_oss_forbid_overwrite"],
        "key": key, "success_action_status": "200",
    }
    with open(path, "rb") as f:
        up = http().post(pol["upload_host"], data=fields, files={"file": (path.name, f)}, timeout=600)
    if up.status_code != 200:
        raise RuntimeError(f"上传到百炼临时存储失败 HTTP {up.status_code}: {up.text[:200]}")
    return f"oss://{key}"


def prepare_media(row):
    """返回交给百炼的地址：普通平台直接用视频直链；B 站或之前下载失败的，先下载再上传。"""
    plat = row["platform"]
    if plat != "bilibili" and not row["force_upload"]:
        return row["media_url"]
    if plat == "bilibili":
        extra = json.loads(row["extra"] or "{}")
        src, ext = bilibili_audio_urls(extra.get("bvid"), extra.get("cid")), ".m4a"
    else:
        src, ext = row["media_url"], ".mp4"
    if not src:
        raise RuntimeError("没有可下载的媒体地址")
    tmp = cfg.OUT_DIR / "_tmp"
    tmp.mkdir(parents=True, exist_ok=True)
    dest = tmp / f"{plat}_{re.sub(r'[^0-9A-Za-z_]', '', row['item_id'] or '')[:40]}_{uuid.uuid4().hex[:8]}{ext}"
    try:
        download(src, plat, dest)
        return upload_to_dashscope(dest)
    finally:
        dest.unlink(missing_ok=True)


def submit_asr(file_urls):
    params = {"language_hints": ["zh", "en"]}
    if cfg.HOTWORDS_VOCAB_ID:
        params["vocabulary_id"] = cfg.HOTWORDS_VOCAB_ID
    body = {"model": cfg.ASR_MODEL, "input": {"file_urls": file_urls}, "parameters": params}
    oss = any(u.startswith("oss://") for u in file_urls)
    for i in range(3):
        r = http().post(f"{cfg.DASHSCOPE_BASE_URL}/api/v1/services/audio/asr/transcription",
                        headers=ds_headers(True, oss), json=body, timeout=60)
        if r.status_code == 429:
            time.sleep(5 * (i + 1))
            continue
        if r.status_code != 200:
            raise RuntimeError(f"提交转写失败 HTTP {r.status_code}: {r.text[:300]}")
        return r.json()["output"]["task_id"]
    raise RuntimeError("提交转写失败：限流")


class TaskGone(Exception):
    pass


def query_asr(task_id):
    r = http().get(f"{cfg.DASHSCOPE_BASE_URL}/api/v1/tasks/{task_id}", headers=ds_headers(), timeout=60)
    if r.status_code in (400, 404):
        raise TaskGone(f"HTTP {r.status_code}")
    r.raise_for_status()
    return r.json().get("output") or {}


def fetch_transcript(transcription_url):
    r = http().get(transcription_url, timeout=60)
    r.raise_for_status()
    j = r.json()
    texts, sentences = [], []
    for t in j.get("transcripts") or []:
        if t.get("text"):
            texts.append(t["text"])
        for s in t.get("sentences") or []:
            sentences.append({"start": s.get("begin_time"), "end": s.get("end_time"), "text": s.get("text", "")})
    dur_ms = (j.get("properties") or {}).get("original_duration_in_milliseconds")
    return "\n".join(texts).strip(), sentences, dur_ms


def match_result(results, asr_url, n_items):
    """按提交地址找结果。oss:// 地址在结果里可能被换成 https 地址，按对象路径兜底匹配。"""
    by_url = {r.get("file_url"): r for r in results}
    if asr_url in by_url:
        return by_url[asr_url]
    if asr_url and asr_url.startswith("oss://"):
        key = asr_url[len("oss://"):]
        for fu, r in by_url.items():
            if fu and key in unquote(fu):
                return r
    if n_items == 1 and len(results) == 1:
        return results[0]
    return None


# ----------------------------------------------------------------------------
# 输出
# ----------------------------------------------------------------------------
def safe_name(s, n=40):
    s = re.sub(r"[\\/:*?\"<>|#@\x00-\x1f\x7f]+", " ", s or "").strip()
    return re.sub(r"\s+", "_", s)[:n].rstrip(". ") or "untitled"


def ms_to_ts(ms):
    if ms is None:
        return ""
    s = int(ms) // 1000
    return f"{s // 3600}:{s // 60 % 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60:02d}:{s % 60:02d}"


def paragraphs(sentences, gap_ms=1500, max_chars=200):
    """把句子按停顿和长度合并成段落，阅读更舒服。"""
    paras, cur, cur_start, last_end = [], [], None, None
    for s in sentences:
        if cur and ((s["start"] or 0) - (last_end or 0) > gap_ms or sum(len(x) for x in cur) > max_chars):
            paras.append((cur_start, "".join(cur)))
            cur, cur_start = [], None
        if cur_start is None:
            cur_start = s["start"]
        cur.append(s["text"])
        last_end = s["end"]
    if cur:
        paras.append((cur_start, "".join(cur)))
    return paras


def fmt_n(v):
    return "—" if v is None else f"{v:,}"


def stats_line(row):
    parts = [("播放", row["plays"]), ("点赞", row["likes"]), ("评论", row["comments"]),
             ("收藏", row["collects"]), ("分享", row["shares"])]
    if row["platform"] == "bilibili":
        parts += [("投币", row["coins"]), ("弹幕", row["danmaku"])]
    return " · ".join(f"{k} {fmt_n(v)}" for k, v in parts)


def write_markdown(conn, row):
    d = cfg.OUT_DIR / "transcripts"
    d.mkdir(parents=True, exist_ok=True)
    ident = row["item_id"] or hashlib.md5(row["url"].encode()).hexdigest()[:10]
    path = d / f"{row['platform']}_{ident}_{safe_name(row['title'])}.md"
    sentences = json.loads(row["sentences"] or "[]")
    if row["status"] == "done":
        body = "\n\n".join(f"`{ms_to_ts(st)}` {t}" for st, t in paragraphs(sentences)) if sentences else (row["text"] or "")
        body = body or "（未识别到人声）"
    else:
        body = "（未转写，本次只获取了数据）"
    lines = [
        f"# {row['title'] or '(无标题)'}", "",
        f"- 平台：{PLATFORM_ZH.get(row['platform'], row['platform'])}",
        f"- 作者：{row['author'] or ''}" + (f"（ID {row['author_id']}）" if row["author_id"] else ""),
        f"- 发布时间：{row['publish_time'] or '—'}",
        f"- 时长：{row['duration_sec'] or '—'} 秒",
        f"- 数据：{stats_line(row)}（{(row['stats_at'] or '')[:16].replace('T', ' ')} 获取）",
    ]
    if row["tags"]:
        lines.append(f"- 话题：{row['tags']}")
    lines += [f"- 链接：{row['url']}", "", "## 逐字稿", "", body, ""]
    comments = json.loads(row["top_comments"] or "null")
    if comments:
        lines += ["## 热门评论", ""]
        for i, c in enumerate(comments, 1):
            meta = "，".join(x for x in (f"赞 {c['likes']:,}" if c.get("likes") is not None else "", c.get("time") or "") if x)
            lines.append(f"{i}. **{c.get('user') or '匿名'}**{f'（{meta}）' if meta else ''}：{c.get('text', '')}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    upsert(conn, row["url"], md_path=str(path))
    return path


EXPORT_COLS = [("platform", "平台"), ("url", "链接"), ("status", "状态"), ("item_id", "作品ID"), ("title", "标题"),
               ("author", "作者"), ("publish_time", "发布时间"), ("duration_sec", "时长(秒)"), ("plays", "播放"),
               ("likes", "点赞"), ("comments", "评论数"), ("collects", "收藏"), ("shares", "分享"), ("coins", "投币"),
               ("danmaku", "弹幕"), ("tags", "话题"), ("stats_at", "数据获取时间"), ("text", "逐字稿"),
               ("error", "错误信息")]


def export_values(r):
    vals = []
    for c, _ in EXPORT_COLS:
        v = r[c]
        if c == "platform":
            v = PLATFORM_ZH.get(v, v)
        elif c == "status":
            v = STATUS_ZH.get(v, v)
        vals.append("" if v is None else v)
    return vals


def export(conn):
    rows = conn.execute("SELECT * FROM items ORDER BY rowid").fetchall()
    heads = [h for _, h in EXPORT_COLS]
    csv_path = cfg.OUT_DIR / "summary.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig：Excel 直接打开不乱码
        w = csv.writer(f)
        w.writerow(heads)
        for r in rows:
            w.writerow(export_values(r))
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font
    except ImportError:
        log("未安装 openpyxl，跳过 summary.xlsx（CSV 已生成）")
        return csv_path

    def clean(v):  # 标题/逐字稿里的控制字符会让 openpyxl 报错
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:32000] if isinstance(v, str) else v

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(heads)
    for r in rows:
        ws.append([clean(v) for v in export_values(r)])
    widths = {"链接": 40, "标题": 40, "逐字稿": 80, "错误信息": 30, "话题": 24, "作品ID": 22, "发布时间": 17,
              "数据获取时间": 20}
    text_col = heads.index("逐字稿")
    for i, h in enumerate(heads):
        ws.column_dimensions[chr(65 + i)].width = widths.get(h, 10)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=text_col + 1, max_col=text_col + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    crow = [(r, json.loads(r["top_comments"] or "[]")) for r in rows]
    if any(cs for _, cs in crow):
        ws2 = wb.create_sheet("评论")
        ws2.append(["平台", "标题", "链接", "序号", "用户", "点赞", "时间", "评论内容"])
        for c in ws2[1]:
            c.font = Font(bold=True)
        for r, cs in crow:
            for i, c in enumerate(cs, 1):
                ws2.append([clean(x) for x in (PLATFORM_ZH.get(r["platform"]), r["title"], r["url"], i,
                                               c.get("user"), c.get("likes"), c.get("time"), c.get("text"))])
        for col, wdt in zip("ABCDEFGH", (8, 30, 30, 6, 16, 8, 17, 80)):
            ws2.column_dimensions[col].width = wdt
    wb.save(cfg.OUT_DIR / "summary.xlsx")
    return csv_path


# ----------------------------------------------------------------------------
# 主流程
# ----------------------------------------------------------------------------
DOWNLOAD_HINTS = ("DOWNLOAD", "FILE", "403", "FORBIDDEN")


def fail_task_items(conn, tid, err):
    for it in conn.execute("SELECT url FROM items WHERE task_id=? AND status='submitted'", (tid,)).fetchall():
        upsert(conn, it["url"], status="failed", error=err)


def resubmit_stuck_task(conn, tid):
    """百炼偶尔会让一个多文件任务卡住（单独提交同样的文件几十秒就完成）。
    超时后把任务里的文件重新提交一次；已经重试过的仍超时，才记为失败。"""
    items = conn.execute("SELECT url, asr_url, asr_tries FROM items WHERE task_id=? AND status='submitted'",
                         (tid,)).fetchall()
    retry = [it for it in items if (it["asr_tries"] or 0) < 2]
    for it in items:
        if it not in retry:
            upsert(conn, it["url"], status="failed",
                   error=f"转写两次都超过 {cfg.TASK_TIMEOUT_MIN} 分钟未完成，可用 --retry-failed 重跑")
    if not retry:
        return
    try:
        new_tid = submit_asr(list(dict.fromkeys(it["asr_url"] for it in retry)))
    except Exception as e:
        log(f"重新提交卡住的任务 {tid[:8]} 失败，稍后再试：{e}")
        return
    for it in retry:
        upsert(conn, it["url"], task_id=new_tid, submitted_at=now(), asr_tries=(it["asr_tries"] or 0) + 1)
    log(f"转写任务 {tid[:8]} 超过 {cfg.TASK_TIMEOUT_MIN} 分钟未完成，已把 {len(retry)} 个文件重新提交为 {new_tid[:8]}")


def poll_inflight(conn, block_until=None):
    """查询进行中的转写任务；block_until=N 时一直等到进行中的任务数 <= N。"""
    while True:
        tasks = conn.execute("SELECT task_id, MIN(submitted_at) AS t0 FROM items "
                             "WHERE status='submitted' AND task_id IS NOT NULL GROUP BY task_id").fetchall()
        for t in tasks:
            tid = t["task_id"]
            try:
                out = query_asr(tid)
            except TaskGone as e:
                fail_task_items(conn, tid, f"转写任务已失效（{e}），可用 --retry-failed 重跑")
                continue
            except Exception as e:
                log(f"查询任务 {tid[:8]} 出错，稍后重试：{e}")
                continue
            st = out.get("task_status")
            if st in ("PENDING", "RUNNING"):
                started = datetime.fromisoformat(t["t0"]) if t["t0"] else datetime.now()
                if (datetime.now() - started).total_seconds() > cfg.TASK_TIMEOUT_MIN * 60:
                    resubmit_stuck_task(conn, tid)
                continue
            if st not in ("SUCCEEDED", "FAILED"):
                # UNKNOWN 等：任务过期（百炼只保留 24 小时）或不存在，不能无限等下去
                fail_task_items(conn, tid, f"转写任务状态为 {st}（可能已过期），可用 --retry-failed 重跑")
                continue
            results = out.get("results") or []
            items = conn.execute("SELECT * FROM items WHERE task_id=? AND status='submitted'", (tid,)).fetchall()
            ok = bad = 0
            for it in items:
                res = match_result(results, it["asr_url"], len(items))
                if res and res.get("subtask_status") == "SUCCEEDED" and res.get("transcription_url"):
                    try:
                        text, sents, dur_ms = fetch_transcript(res["transcription_url"])
                        fields = dict(status="done", text=text, sentences=json.dumps(sents, ensure_ascii=False),
                                      error=None)
                        if dur_ms:
                            fields["duration_sec"] = round(dur_ms / 1000, 1)
                        upsert(conn, it["url"], **fields)
                        ok += 1
                    except Exception as e:
                        err, force = f"下载转写结果失败：{e}", 0
                    else:
                        try:
                            write_markdown(conn, get_row(conn, it["url"]))
                        except OSError as e:  # 逐字稿已存进状态库，写文件失败不算转写失败
                            log(f"写入 Markdown 失败（{it['url']}）：{e}")
                        continue
                else:
                    err = str((res or {}).get("message") or (res or {}).get("code") or out.get("message") or "转写失败")
                    force = int(any(h in err.upper() for h in DOWNLOAD_HINTS))
                    if force:
                        err += "（视频地址过期或有防盗链；--retry-failed 会改为先下载再上传）"
                upsert(conn, it["url"], status="failed", error=err[:500], force_upload=force or it["force_upload"])
                bad += 1
            log(f"转写任务 {tid[:8]} 完成：成功 {ok}，失败 {bad}")
        n = conn.execute("SELECT COUNT(DISTINCT task_id) FROM items WHERE status='submitted' "
                         "AND task_id IS NOT NULL").fetchone()[0]
        if block_until is None or n <= block_until:
            return n
        time.sleep(cfg.POLL_INTERVAL)


def run_parallel(fn, keys):
    """并发执行 fn(key)，按完成顺序产出 (key, 结果, 异常)。数据库只在主线程写。"""
    with ThreadPoolExecutor(max_workers=max(1, cfg.PARSE_WORKERS)) as ex:
        futs = {ex.submit(fn, k): k for k in keys}
        for f in as_completed(futs):
            try:
                yield futs[f], f.result(), None
            except Exception as e:
                yield futs[f], None, e


def parse_chunk(conn, urls, want_media):
    plat_of = {u: get_row(conn, u)["platform"] for u in urls}
    ok = []
    for u, info, err in run_parallel(lambda u: PARSERS[plat_of[u]](u, want_media), urls):
        if err:
            conn.execute("UPDATE items SET attempts=attempts+1 WHERE url=?", (u,))
            upsert(conn, u, status="failed", error=f"解析失败：{str(err)[:400]}")
            continue
        info["extra"] = json.dumps(info.get("extra") or {}, ensure_ascii=False)
        upsert(conn, u, **info, status="parsed", stats_at=now(), error=None)
        ok.append(u)
    return ok


def prepare_chunk(conn, urls):
    rows = {u: get_row(conn, u) for u in urls}
    ok = []
    for u, asr_url, err in run_parallel(lambda u: prepare_media(rows[u]), urls):
        if err or not asr_url:
            upsert(conn, u, status="failed", error=f"准备音视频失败：{str(err or '没有媒体地址')[:400]}")
        else:
            upsert(conn, u, asr_url=asr_url)
            ok.append((u, asr_url))
    return ok


def enrich(conn, urls, comments_n, douyin_play):
    """可选的额外数据：评论内容、抖音播放量。只补缺失的，不重复扣费。"""
    jobs = []
    for u in urls:
        r = get_row(conn, u)
        # 转写失败但解析成功（有 stats_at）的条目，数据照样补齐
        if not r or not r["stats_at"] or r["status"] == "pending" or not (r["item_id"] or r["platform"] == "xhs"):
            continue
        if douyin_play and r["platform"] == "douyin" and r["plays"] is None:
            jobs.append(("plays", u))
        if comments_n and r["top_comments"] is None:
            jobs.append(("comments", u))
    if not jobs:
        return
    log(f"获取额外数据 {len(jobs)} 项（评论/播放量）…")
    snapshot = {u: dict(get_row(conn, u)) for _, u in jobs}  # 工作线程不碰数据库连接

    def work(job):
        kind, u = job
        r = snapshot[u]
        return douyin_plays(r["item_id"]) if kind == "plays" else fetch_comments(r, comments_n)

    for (kind, u), val, err in run_parallel(work, jobs):
        if err:
            log(f"{'播放量' if kind == 'plays' else '评论'}获取失败（{u}）：{str(err)[:150]}")
        elif kind == "plays":
            upsert(conn, u, plays=val)
        else:
            upsert(conn, u, top_comments=json.dumps(val, ensure_ascii=False))


def row_summary(r):
    keys = ["url", "platform", "status", "item_id", "title", "author", "author_id", "publish_time", "duration_sec",
            "plays", "likes", "comments", "collects", "shares", "coins", "danmaku", "tags", "stats_at", "md_path",
            "error"]
    d = {k: r[k] for k in keys}
    d["transcript_chars"] = len(r["text"] or "")
    d["top_comments_count"] = len(json.loads(r["top_comments"] or "[]"))
    return d


def run(args):
    if not cfg.TIKHUB_API_KEY or (not args.stats_only and not cfg.DASHSCOPE_API_KEY):
        need = "TIKHUB_API_KEY" + ("" if args.stats_only else " 和 DASHSCOPE_API_KEY")
        sys.exit(f"缺少 {need}。请在 {DEFAULT_CONFIG} 里填写（参考技能目录下的 transcript.env.example），"
                 f"或设置同名环境变量。运行 --check 可查看当前配置。")
    conn = db()

    lines = list(args.inputs or [])
    if args.file:
        lines += read_file_lines(Path(args.file))
    links = extract_links(lines)
    if lines and not links:
        log("没有识别到支持的链接（支持抖音、小红书、快手、B 站）")
    for url, plat, raw in links:
        conn.execute("INSERT OR IGNORE INTO items(url, platform, raw_input) VALUES (?,?,?)", (url, plat, raw))
    conn.commit()
    scope = [u for u, _, _ in links] or None  # None = 状态库里的全部条目
    scope_set = set(scope or [])

    if args.retry_failed:
        failed = conn.execute("SELECT url FROM items WHERE status='failed'").fetchall()
        for r in failed:
            if scope is None or r["url"] in scope_set:
                upsert(conn, r["url"], status="pending", task_id=None, asr_url=None, error=None)

    def in_scope(where):
        rows = [r["url"] for r in conn.execute(f"SELECT url FROM items WHERE {where} ORDER BY rowid")]
        return rows if scope is None else [u for u in rows if u in scope_set]

    stats = dict(conn.execute("SELECT status, COUNT(*) FROM items GROUP BY status").fetchall())
    log(f"本次识别 {len(links)} 条链接；状态库：{ {STATUS_ZH.get(k, k): v for k, v in stats.items()} }")

    poll_inflight(conn)  # 先把上次没查完的转写任务收尾

    want_tx = not args.stats_only
    # 需要转写时，之前只拿过数据的条目也重新解析一次（视频直链会过期）
    todo = in_scope("status='pending'" + (" OR status='parsed'" if want_tx else ""))
    if args.limit:
        todo = todo[:args.limit]
    total = len(todo)

    for start in range(0, total, cfg.BATCH_SIZE):
        chunk = todo[start:start + cfg.BATCH_SIZE]
        log(f"解析第 {start + 1}-{start + len(chunk)} / {total} 条 …")
        parsed = parse_chunk(conn, chunk, want_tx)
        if not want_tx or not parsed:
            continue
        prepared = prepare_chunk(conn, parsed)
        if not prepared:
            continue
        poll_inflight(conn, block_until=cfg.MAX_INFLIGHT_TASKS - 1)
        uniq = list(dict.fromkeys(a for _, a in prepared))  # 同一个地址在一个任务里只能出现一次
        try:
            tid = submit_asr(uniq)
            for u, _ in prepared:
                upsert(conn, u, status="submitted", task_id=tid, submitted_at=now(), asr_tries=1)
            log(f"已提交转写任务 {tid[:8]}（{len(uniq)} 个文件）；本批失败 {len(chunk) - len(prepared)} 条")
        except Exception as e:
            for u, _ in prepared:
                upsert(conn, u, status="failed", error=str(e)[:400])
            log(f"提交转写失败：{e}")

    # 转写进行中时顺便抓评论 / 播放量
    if args.comments or args.douyin_plays:
        enrich(conn, in_scope("1=1"), args.comments, args.douyin_plays)

    if want_tx:
        log("等待转写完成 …")
        poll_inflight(conn, block_until=0)

    finished = in_scope("status IN ('done','parsed')")
    for u in finished:
        write_markdown(conn, get_row(conn, u))
    csv_path = export(conn)

    stats = dict(conn.execute("SELECT status, COUNT(*) FROM items GROUP BY status").fetchall())
    log(f"全部完成：{ {STATUS_ZH.get(k, k): v for k, v in stats.items()} }")
    log(f"逐字稿：{cfg.OUT_DIR / 'transcripts'}    汇总表：{csv_path}")
    if stats.get("failed"):
        log("有失败条目：看汇总表“错误信息”列，或运行 --retry-failed 重试")
    return [row_summary(get_row(conn, u)) for u in in_scope("1=1")]


def check():
    import importlib.util
    info = {
        "config_files_loaded": cfg.loaded,
        "default_config_path": str(DEFAULT_CONFIG),
        "TIKHUB_API_KEY": bool(cfg.TIKHUB_API_KEY),
        "DASHSCOPE_API_KEY": bool(cfg.DASHSCOPE_API_KEY),
        "TIKHUB_BASE_URL": cfg.TIKHUB_BASE_URL,
        "ASR_MODEL": cfg.ASR_MODEL,
        "out_dir": str(cfg.OUT_DIR),
        "openpyxl_installed": importlib.util.find_spec("openpyxl") is not None,
        "python": sys.version.split()[0],
    }
    info["ready"] = info["TIKHUB_API_KEY"] and info["DASHSCOPE_API_KEY"]
    return info


def main():
    ap = argparse.ArgumentParser(description="抖音/小红书/快手/B站 视频逐字稿 + 互动数据")
    ap.add_argument("inputs", nargs="*", help="链接或整段分享文案，可以给多个")
    ap.add_argument("--file", help="链接清单文件（.txt 每行一条，或 .csv）")
    ap.add_argument("--out", help="输出目录（默认：当前目录下的 relie-transcripts）")
    ap.add_argument("--config", help=f"配置文件路径（默认：{DEFAULT_CONFIG}）")
    ap.add_argument("--limit", type=int, help="本次最多处理多少条（试跑用）")
    ap.add_argument("--stats-only", action="store_true", help="只获取作品信息和互动数据，不转写")
    ap.add_argument("--comments", type=int, default=0, metavar="N", help="额外抓取前 N 条评论（按页额外计费）")
    ap.add_argument("--douyin-plays", action="store_true", help="额外查询抖音播放量（每条额外计费 1 次）")
    ap.add_argument("--retry-failed", action="store_true", help="重新处理失败的条目")
    ap.add_argument("--export", action="store_true", help="只根据已有结果重新导出汇总表")
    ap.add_argument("--check", action="store_true", help="检查配置和依赖，不调用任何接口")
    ap.add_argument("--json", action="store_true", help="把结果以 JSON 打印到标准输出")
    a = ap.parse_args()
    init_config(a.config, a.out)

    if a.check:
        info = check()
        print(json.dumps(info, ensure_ascii=False, indent=1))
        sys.exit(0 if info["ready"] else 1)
    if a.export:
        conn = db()
        path = export(conn)
        log(f"已导出：{path}")
        if a.json:
            rows = conn.execute("SELECT * FROM items ORDER BY rowid").fetchall()
            print(json.dumps([row_summary(r) for r in rows], ensure_ascii=False, indent=1))
        return
    result = run(a)
    if a.json:
        print(json.dumps({"out_dir": str(cfg.OUT_DIR), "items": result}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
